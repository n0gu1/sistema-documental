"""Export all traceability evidence without the generic report's row limit."""
import json
from math import ceil
from io import BytesIO
from xml.sax.saxutils import escape

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer


def date_text(value):
    return timezone.localtime(value).strftime('%Y-%m-%d %H:%M:%S %z') if value else 'No consta'


def actor_text(actor):
    return f"{actor['name']} ({actor['id']})" if actor else 'No consta'


def result_text(value):
    return 'Exitoso' if value is True else 'Fallido' if value is False else 'No consta'


def event_values(row):
    return [row['code'], row['title'], row['version'] or '', date_text(row['at']), row['event'],
            actor_text(row['actor']), actor_text(row['reviewer']), row['review_id'] or '',
            row['state_from'] or '', row['state_to'] or '',
            row['restored_from']['version'] if row['restored_from'] else '',
            result_text(row['successful']), row['comment'], row['source'], row['source_id'],
            json.dumps(row['details'], ensure_ascii=False, default=str), row['version_id'] or '']


HEADERS = ['Código', 'Documento', 'Versión', 'Fecha', 'Evento', 'Actor', 'Revisor asignado', 'Solicitud',
           'Estado anterior', 'Estado posterior', 'Restaurada desde', 'Resultado', 'Comentario',
           'Fuente', 'ID fuente', 'Detalle', 'ID versión']


def build_traceability_xlsx(data):
    book = Workbook()
    info = book.active
    info.title = 'Documento'
    info.append(['Campo', 'Valor'])
    for key, value in data['document'].items():
        info.append([key, str(value)])
    info.append(['Zona horaria', data['timezone']])
    for note in data['notes']:
        info.append(['Nota', note])
    inventory = book.create_sheet('Versiones')
    inventory.append(['ID', 'Versión', 'Autor', 'Creación', 'Estado actual', 'Vigente', 'Comentario'])
    for version in data['versions']:
        inventory.append([version['id'], version['version'], actor_text(version['author']), date_text(version['created_at']),
                          version['current_state'], 'Sí' if version['is_current'] else 'No', version['comment']])
    sheet = book.create_sheet('Cronología')
    sheet.append(HEADERS)
    for row in data['rows']:
        sheet.append(event_values(row))
    for tab in book:
        tab.freeze_panes = 'A2'
        tab.auto_filter.ref = tab.dimensions
        for row in tab:
            for cell in row:
                cell.data_type = 's'  # Literal evidence, never executable formulas.
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        for cell in tab[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F4E78')
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = min(60, max(16, max(len(str(c.value or '')) for c in column) + 2))
        for row in tab:
            lines = max(sum(max(1, ceil(len(line) / (tab.column_dimensions[c.column_letter].width - 2)))
                            for line in str(c.value or '').split('\n')) for c in row)
            tab.row_dimensions[row[0].row].height = min(409, max(24, 15 * lines + 6))
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def build_traceability_pdf(data):
    output = BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    story = []

    def para(text, style='Normal'):
        story.append(Paragraph(escape(str(text)).replace('\n', '<br/>'), styles[style]))

    para('Reporte integral de trazabilidad', 'Title')
    para(f"{data['document']['code']} — {data['document']['title']}", 'Heading2')
    para(f"Documento: {data['document']['id']} | Zona horaria: {data['timezone']}")
    for note in data['notes']:
        para(note)
    para('Versiones conservadas', 'Heading2')
    for v in data['versions']:
        para(f"Versión {v['version']} | Estado actual: {v['current_state']} | Vigente: {'Sí' if v['is_current'] else 'No'}", 'Heading3')
        para(f"ID: {v['id']} | Creación: {date_text(v['created_at'])}")
        para(f"Autor: {actor_text(v['author'])}")
        para(v['comment'] or 'Sin comentario')
    para(f"Cronología: {len(data['rows'])} registros", 'Heading2')
    for row in data['rows']:
        start = len(story)
        para(f"{date_text(row['at'])} — {row['event']}", 'Heading3')
        para(f"Fuente: {row['source']} / {row['source_id']} | Código: {row['event_code']}")
        if row['version']:
            para(f"Versión: {row['version']} ({row['version_id']})")
        para(f"Actor: {actor_text(row['actor'])} | Resultado: {result_text(row['successful'])}")
        if row['review_id']:
            para(f"Solicitud: {row['review_id']} | Revisor asignado: {actor_text(row['reviewer'])}")
        if row['state_from'] or row['state_to']:
            para(f"Estado: {row['state_from'] or 'No consta'} → {row['state_to'] or 'No consta'}")
        if row['restored_from']:
            para(f"Restaurada desde: {row['restored_from']['version']} ({row['restored_from']['id']})")
        if row['comment']:
            para(row['comment'])
        if row['details']:
            para(json.dumps(row['details'], ensure_ascii=False, default=str))
        story.append(Spacer(1, 6))
        story[start:] = [KeepTogether(story[start:])]

    def page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.drawRightString(576, 20, f'Página {doc.page}')
        canvas.restoreState()
    pdf.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return output.getvalue()
